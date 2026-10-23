import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule

FECHA_CORTE = pd.Timestamp(pd.Timestamp.today())
ARCHIVO_ENTRADA = "facturas_prueba.csv"
ARCHIVO_SALIDA = "reporte_cobranzas.xlsx"


def limpiar_monto(val):
    if pd.isna(val) or str(val).strip() == "":
        return np.nan
    s = str(val).strip().replace("$", "").strip()
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def parsear_fecha(val):
    if pd.isna(val):
        return pd.NaT
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return pd.to_datetime(s, format=fmt)
        except ValueError:
            continue
    return pd.to_datetime(s, errors="coerce")


def normalizar_cliente(nombre):
    mapeo_siglas = {"S.a.": "S.A.", "Srl": "SRL", "Gba": "GBA", "Ba": "BA", "Del": "del"}
    limpio = " ".join(str(nombre).strip().split()).title()
    palabras = [mapeo_siglas.get(p, p) for p in limpio.split(" ")]
    return " ".join(palabras)


def normalizar_estado(est):
    e = str(est).strip().upper()
    return "Pagado" if e in ["PAGADO", "PAGADA", "COBRADO"] else "Pendiente"


def procesar_datos(ruta_csv):
    df = pd.read_csv(ruta_csv)

    df["id_factura"] = df["id_factura"].astype(str).str.strip().str.upper()
    df["cliente_limpio"] = df["cliente"].apply(normalizar_cliente)
    df["monto_num"] = df["monto"].apply(limpiar_monto)
    df["fecha_venc_dt"] = df["fecha_vencimiento"].apply(parsear_fecha)
    df["estado_limpio"] = df["estado_pago"].apply(normalizar_estado)

    auditoria = []
    indices_descartar = set()

    dup_id = df.duplicated(subset=["id_factura"], keep="first")
    for idx, row in df[dup_id].iterrows():
        auditoria.append({
            "id_factura": row["id_factura"],
            "cliente": row["cliente_limpio"],
            "monto_original": row["monto"],
            "monto_numerico": row["monto_num"],
            "fecha_vencimiento": row["fecha_venc_dt"].strftime("%Y-%m-%d") if pd.notna(row["fecha_venc_dt"]) else str(row["fecha_vencimiento"]),
            "estado_pago": row["estado_limpio"],
            "motivo_alerta": "Duplicado de ID de factura"
        })
        indices_descartar.add(idx)

    df_sin_dup_id = df.drop(index=list(indices_descartar))
    dup_negocio = df_sin_dup_id.duplicated(subset=["cliente_limpio", "monto_num", "fecha_venc_dt"], keep="first")
    for idx, row in df_sin_dup_id[dup_negocio].iterrows():
        auditoria.append({
            "id_factura": row["id_factura"],
            "cliente": row["cliente_limpio"],
            "monto_original": row["monto"],
            "monto_numerico": row["monto_num"],
            "fecha_vencimiento": row["fecha_venc_dt"].strftime("%Y-%m-%d") if pd.notna(row["fecha_venc_dt"]) else str(row["fecha_vencimiento"]),
            "estado_pago": row["estado_limpio"],
            "motivo_alerta": "Posible duplicado (mismo cliente, monto y vencimiento)"
        })
        indices_descartar.add(idx)

    df_restante = df.drop(index=list(indices_descartar))
    montos_validos = df_restante[df_restante["monto_num"] > 1000]["monto_num"]
    q1, q3 = montos_validos.quantile(0.25), montos_validos.quantile(0.75)
    umbral_outlier = q3 + 3.0 * (q3 - q1)

    for idx, row in df_restante.iterrows():
        m = row["monto_num"]
        motivo = None
        if pd.isna(m):
            motivo = "Monto faltante o ilegible"
        elif m < 0:
            motivo = "Monto negativo (revisar si es Nota de Crédito mal cargada)"
        elif m == 0:
            motivo = "Monto igual a cero ($0.00)"
        elif m < 1000:
            motivo = "Monto atípicamente bajo (< $1.000)"
        elif m > umbral_outlier:
            motivo = f"Outlier extremo (> ${umbral_outlier:,.2f})"

        if motivo:
            auditoria.append({
                "id_factura": row["id_factura"],
                "cliente": row["cliente_limpio"],
                "monto_original": row["monto"],
                "monto_numerico": m,
                "fecha_vencimiento": row["fecha_venc_dt"].strftime("%Y-%m-%d") if pd.notna(row["fecha_venc_dt"]) else str(row["fecha_vencimiento"]),
                "estado_pago": row["estado_limpio"],
                "motivo_alerta": motivo
            })
            indices_descartar.add(idx)

    df_auditoria = pd.DataFrame(auditoria)
    df_limpio = df.drop(index=list(indices_descartar)).copy()


    df_limpio["dias_vencido"] = (FECHA_CORTE - df_limpio["fecha_venc_dt"]).dt.days
    df_limpio["dias_mora"] = np.where(
        (df_limpio["estado_limpio"] == "Pendiente") & (df_limpio["dias_vencido"] > 0),
        df_limpio["dias_vencido"],
        0
    )

    df_limpio["situacion"] = np.where(
        df_limpio["estado_limpio"] == "Pagado", "Pagado",
        np.where(df_limpio["dias_vencido"] > 0, "Vencido", "Pendiente (No vencido)")
    )

    df_limpio["tramo_mora"] = pd.cut(
        df_limpio["dias_mora"],
        bins=[-1, 0, 30, 60, np.inf],
        labels=["Al día / Pagado", "1 a 30 días", "31 a 60 días", "+60 días (Crítico)"]
    )


    df_limpio["monto_vencido"] = np.where(df_limpio["situacion"] == "Vencido", df_limpio["monto_num"], 0.0)
    df_limpio["monto_no_vencido"] = np.where(df_limpio["situacion"] == "Pendiente (No vencido)", df_limpio["monto_num"], 0.0)
    df_limpio["monto_pagado"] = np.where(df_limpio["situacion"] == "Pagado", df_limpio["monto_num"], 0.0)
    df_limpio["es_vencida"] = np.where(df_limpio["situacion"] == "Vencido", 1, 0)

    resumen_clientes = (
        df_limpio.groupby("cliente_limpio", as_index=False)
        .agg(
            cant_facturas=("id_factura", "count"),
            facturas_vencidas=("es_vencida", "sum"),
            total_facturado=("monto_num", "sum"),
            total_pagado=("monto_pagado", "sum"),
            pendiente_no_vencido=("monto_no_vencido", "sum"),
            deuda_vencida=("monto_vencido", "sum"),
            max_dias_mora=("dias_mora", "max")
        )
        .sort_values(by="deuda_vencida", ascending=False)
    )
    resumen_clientes["pct_deuda_sobre_total"] = (
        resumen_clientes["deuda_vencida"] / resumen_clientes["deuda_vencida"].sum()
    )

    return df, df_limpio, df_auditoria, resumen_clientes


def exportar_excel_formateado(df_raw, df_limpio, df_auditoria, resumen_clientes, ruta_salida):
    wb = openpyxl.Workbook()

    COLOR_HEADER_BG = "172A2B"
    COLOR_HEADER_FG = "FFFFFF"
    COLOR_ZEBRA = "F4F7F7"
    COLOR_BORDER = "D1D9D9"
    COLOR_ALERT_BG = "FADBD8"
    COLOR_ALERT_FG = "78281F"
    COLOR_WARN_BG = "FCF3CF"
    COLOR_WARN_FG = "7E5109"
    COLOR_OK_BG = "D4EFDF"
    COLOR_OK_FG = "145A32"

    font_title = Font(name="Calibri", size=15, bold=True, color="172A2B")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="566573")
    font_header = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_FG)
    font_bold = Font(name="Calibri", size=11, bold=True, color="172A2B")
    font_body = Font(name="Calibri", size=11, color="1C2833")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="566573")
    font_kpi_val = Font(name="Calibri", size=14, bold=True, color="172A2B")

    fill_header = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    fill_zebra = PatternFill("solid", fgColor=COLOR_ZEBRA)
    fill_total = PatternFill("solid", fgColor="E5ECEC")
    fill_kpi = PatternFill("solid", fgColor="EBF1F1")

    thin_side = Side(border_style="thin", color=COLOR_BORDER)
    med_bottom = Side(border_style="medium", color="172A2B")
    double_bottom = Side(border_style="double", color="172A2B")

    border_data = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_total = Border(top=thin_side, bottom=double_bottom, left=thin_side, right=thin_side)
    border_kpi = Border(left=thin_side, right=thin_side, top=thin_side, bottom=med_bottom)

  
    ws1 = wb.active
    ws1.title = "Resumen Deuda Clientes"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "REPORTE EJECUTIVO DE COBRANZAS Y DEUDA VENCIDA"
    ws1["A1"].font = font_title
    ws1["A2"] = (
        f"Fecha de corte del análisis: {FECHA_CORTE.strftime('%d/%m/%Y')} | "
        f"Registros auditados: {len(df_raw)} | Anomalías aisladas: {len(df_auditoria)}"
    )
    ws1["A2"].font = font_subtitle

    
    kpis = [
        ("A", "B", "DEUDA VENCIDA TOTAL", resumen_clientes["deuda_vencida"].sum(), '"$"#,##0.00'),
        ("C", "D", "PENDIENTE NO VENCIDO", resumen_clientes["pendiente_no_vencido"].sum(), '"$"#,##0.00'),
        ("E", "F", "TOTAL COBRADO", resumen_clientes["total_pagado"].sum(), '"$"#,##0.00'),
        ("G", "H", "ALERTAS / DUPLICADOS", len(df_auditoria), '#,##0')
    ]
    for col1, col2, label, val, num_fmt in kpis:
        ws1.merge_cells(f"{col1}4:{col2}4")
        ws1.merge_cells(f"{col1}5:{col2}5")
        c_lbl = ws1[f"{col1}4"]
        c_val = ws1[f"{col1}5"]
        c_lbl.value = label
        c_lbl.font = font_kpi_label
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_lbl.fill = fill_kpi
        c_val.value = val
        c_val.font = font_kpi_val
        c_val.number_format = num_fmt
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        c_val.fill = fill_kpi
        for r in [4, 5]:
            for c in [col1, col2]:
                ws1[f"{c}{r}"].border = border_kpi


    headers_ws1 = [
        "Cliente", "Facturas Totales", "Facturas Vencidas", "Total Facturado",
        "Total Pagado", "Pendiente (En Fecha)", "Deuda Vencida Exigible",
        "% Share Deuda", "Máx. Días Mora"
    ]
    start_row_ws1 = 7
    for col_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=start_row_ws1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center", wrap_text=True)
    ws1.row_dimensions[start_row_ws1].height = 26


    for idx, row in enumerate(resumen_clientes.itertuples(index=False), start=start_row_ws1 + 1):
        vals = [
            row.cliente_limpio,
            int(row.cant_facturas),
            int(row.facturas_vencidas),
            float(row.total_facturado),
            float(row.total_pagado),
            float(row.pendiente_no_vencido),
            float(row.deuda_vencida),
            float(row.pct_deuda_sobre_total),
            int(row.max_dias_mora)
        ]
        for col_idx, v in enumerate(vals, 1):
            c = ws1.cell(row=idx, column=col_idx, value=v)
            c.font = font_body
            c.border = border_data
            if idx % 2 == 1:
                c.fill = fill_zebra
            if col_idx == 1:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in [2, 3, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.number_format = "#,##0"
            elif col_idx in [4, 5, 6, 7]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '"$"#,##0.00'
                if col_idx == 7 and v > 0:
                    c.font = Font(name="Calibri", size=11, bold=True, color="78281F")
            elif col_idx == 8:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "0.0%"


    total_row_ws1 = start_row_ws1 + len(resumen_clientes) + 1
    last_data_row = total_row_ws1 - 1

    ws1.cell(row=total_row_ws1, column=1, value="TOTAL GENERAL").alignment = Alignment(horizontal="left", vertical="center")
    ws1.cell(row=total_row_ws1, column=2, value=f"=SUM(B{start_row_ws1+1}:B{last_data_row})").number_format = "#,##0"
    ws1.cell(row=total_row_ws1, column=3, value=f"=SUM(C{start_row_ws1+1}:C{last_data_row})").number_format = "#,##0"
    ws1.cell(row=total_row_ws1, column=4, value=f"=SUM(D{start_row_ws1+1}:D{last_data_row})").number_format = '"$"#,##0.00'
    ws1.cell(row=total_row_ws1, column=5, value=f"=SUM(E{start_row_ws1+1}:E{last_data_row})").number_format = '"$"#,##0.00'
    ws1.cell(row=total_row_ws1, column=6, value=f"=SUM(F{start_row_ws1+1}:F{last_data_row})").number_format = '"$"#,##0.00'
    ws1.cell(row=total_row_ws1, column=7, value=f"=SUM(G{start_row_ws1+1}:G{last_data_row})").number_format = '"$"#,##0.00'
    ws1.cell(row=total_row_ws1, column=8, value=f"=SUM(H{start_row_ws1+1}:H{last_data_row})").number_format = "0.0%"
    ws1.cell(row=total_row_ws1, column=9, value=f"=MAX(I{start_row_ws1+1}:I{last_data_row})").number_format = "#,##0"

    for col_idx in range(1, 10):
        c = ws1.cell(row=total_row_ws1, column=col_idx)
        c.font = font_bold
        c.fill = fill_total
        c.border = border_total
        if col_idx in [2, 3, 9]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [4, 5, 6, 7, 8]:
            c.alignment = Alignment(horizontal="right", vertical="center")

    rule_high_debt = CellIsRule(
        operator="greaterThan",
        formula=["500000"],
        stopIfTrue=True,
        fill=PatternFill("solid", fgColor=COLOR_ALERT_BG),
        font=Font(name="Calibri", size=11, bold=True, color=COLOR_ALERT_FG)
    )
    ws1.conditional_formatting.add(f"G{start_row_ws1+1}:G{last_data_row}", rule_high_debt)

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Deuda Vencida Exigible por Cliente ($)"
    chart.y_axis.title = "Monto ($)"
    chart.x_axis.title = "Cliente"
    chart.legend = None
    chart.width = 20
    chart.height = 11

    data_ref = Reference(ws1, min_col=7, min_row=start_row_ws1, max_row=last_data_row)
    cats_ref = Reference(ws1, min_col=1, min_row=start_row_ws1 + 1, max_row=last_data_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws1.add_chart(chart, f"A{total_row_ws1 + 2}")

    ws1.freeze_panes = f"A{start_row_ws1 + 1}"
    ws1.auto_filter.ref = f"A{start_row_ws1}:I{last_data_row}"

    ws2 = wb.create_sheet(title="Facturas Limpias")
    ws2["A1"] = "DETALLE DE FACTURAS VALIDADAS Y NORMALIZADAS"
    ws2["A1"].font = font_title

    headers_ws2 = [
        "ID Factura", "Cliente Normalizado", "Monto Validado", "Fecha Vencimiento",
        "Estado Declarado", "Situación Real", "Días de Mora", "Tramo de Antigüedad"
    ]
    start_row_ws2 = 3
    for col_idx, h in enumerate(headers_ws2, 1):
        c = ws2.cell(row=start_row_ws2, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx != 2 else "left", vertical="center")
    ws2.row_dimensions[start_row_ws2].height = 24

    df_limpio_sorted = df_limpio.sort_values(by=["dias_mora", "monto_num"], ascending=[False, False])
    for idx, row in enumerate(df_limpio_sorted.itertuples(index=False), start=start_row_ws2 + 1):
        vals = [
            row.id_factura,
            row.cliente_limpio,
            float(row.monto_num),
            row.fecha_venc_dt.strftime("%Y-%m-%d"),
            row.estado_limpio,
            row.situacion,
            int(row.dias_mora),
            str(row.tramo_mora)
        ]
        for col_idx, v in enumerate(vals, 1):
            c = ws2.cell(row=idx, column=col_idx, value=v)
            c.font = font_body
            c.border = border_data
            if idx % 2 == 1:
                c.fill = fill_zebra
            if col_idx in [1, 4, 5, 6, 7, 8]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 3:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '"$"#,##0.00'

            if col_idx == 6:
                if v == "Vencido":
                    c.fill = PatternFill("solid", fgColor=COLOR_ALERT_BG)
                    c.font = Font(name="Calibri", size=11, bold=True, color=COLOR_ALERT_FG)
                elif v == "Pagado":
                    c.fill = PatternFill("solid", fgColor=COLOR_OK_BG)
                    c.font = Font(name="Calibri", size=11, bold=True, color=COLOR_OK_FG)
                else:
                    c.fill = PatternFill("solid", fgColor=COLOR_WARN_BG)
                    c.font = Font(name="Calibri", size=11, bold=True, color=COLOR_WARN_FG)

    ws2.freeze_panes = f"A{start_row_ws2 + 1}"
    ws2.auto_filter.ref = f"A{start_row_ws2}:H{start_row_ws2 + len(df_limpio_sorted)}"


    ws3 = wb.create_sheet(title="Auditoria Anomalias")
    ws3["A1"] = "REGISTROS SEPARADOS POR DUPLICIDAD O ANOMALÍAS EN MONTOS"
    ws3["A1"].font = font_title
    ws3["A2"] = "Estos registros fueron excluidos del cálculo automático de deuda para evitar distorsiones contables."
    ws3["A2"].font = font_subtitle

    headers_ws3 = [
        "ID Factura", "Cliente Detectado", "Valor Crudo Original", "Monto Interpretado",
        "Fecha Vencimiento", "Estado Pago", "Motivo de Exclusión / Alerta"
    ]
    start_row_ws3 = 4
    for col_idx, h in enumerate(headers_ws3, 1):
        c = ws3.cell(row=start_row_ws3, column=col_idx, value=h)
        c.font = font_header
        c.fill = PatternFill("solid", fgColor="78281F")
        c.alignment = Alignment(horizontal="center" if col_idx not in [2, 7] else "left", vertical="center")
    ws3.row_dimensions[start_row_ws3].height = 24

    for idx, row in enumerate(df_auditoria.itertuples(index=False), start=start_row_ws3 + 1):
        monto_interp = float(row.monto_numerico) if pd.notna(row.monto_numerico) else "N/A"
        vals = [
            row.id_factura,
            row.cliente,
            str(row.monto_original) if pd.notna(row.monto_original) else "(Vacío)",
            monto_interp,
            row.fecha_vencimiento,
            row.estado_pago,
            row.motivo_alerta
        ]
        for col_idx, v in enumerate(vals, 1):
            c = ws3.cell(row=idx, column=col_idx, value=v)
            c.font = font_body
            c.border = border_data
            if idx % 2 == 1:
                c.fill = PatternFill("solid", fgColor="FDEDEC")
            if col_idx in [1, 3, 5, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 7]:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 4:
                c.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(v, float):
                    c.number_format = '"$"#,##0.00'
            if col_idx == 7:
                c.font = Font(name="Calibri", size=11, bold=True, color=COLOR_ALERT_FG)

    ws3.freeze_panes = f"A{start_row_ws3 + 1}"
    ws3.auto_filter.ref = f"A{start_row_ws3}:G{start_row_ws3 + len(df_auditoria)}"

    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 3:
                    continue
                val_str = str(cell.value or "")
                if val_str.startswith("="):
                    val_str = "$ 9.999.999,00"
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    wb.save(ruta_salida)


if __name__ == "__main__":
    df_raw, df_limpio, df_auditoria, resumen = procesar_datos(ARCHIVO_ENTRADA)
    exportar_excel_formateado(df_raw, df_limpio, df_auditoria, resumen, ARCHIVO_SALIDA)
    print(f"Reporte completo generado exitosamente en: {ARCHIVO_SALIDA}")